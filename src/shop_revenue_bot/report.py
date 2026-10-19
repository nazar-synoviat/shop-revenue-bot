"""Amount parsing, text summaries and the Excel export."""

import re
from collections import defaultdict
from datetime import date
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .db import Sale

# "1250", "1 250,50", "1250.5 кава з собою"; a space only groups thousands,
# so "350 2 штуки" stays 350 with the note "2 штуки"
AMOUNT_RE = re.compile(
    r"^\s*((?:\d{1,3}(?:[  ]\d{3})+|\d+)(?:[.,]\d{1,2})?)\s*(?:грн|₴)?\s*(.*)$",
    re.IGNORECASE,
)
WEEKDAYS = ["Пн", "Вт", "Ср", "Чт", "Пт", "Сб", "Нд"]


def parse_sale(text: str) -> tuple[float, str] | None:
    m = AMOUNT_RE.match(text)
    if not m:
        return None
    amount = float(re.sub(r"[  ]", "", m.group(1)).replace(",", "."))
    if amount <= 0:
        return None
    return amount, m.group(2).strip()


def money(value: float) -> str:
    """1234.5 -> '1 234,50 ₴'"""
    return f"{value:,.2f}".replace(",", " ").replace(".", ",") + " ₴"


def day_summary(sales: list[Sale]) -> str:
    if not sales:
        return "Продажів ще немає."
    total = sum(s.amount for s in sales)
    return (
        f"Разом: <b>{money(total)}</b>\n"
        f"Продажів: {len(sales)}\n"
        f"Середній чек: {money(total / len(sales))}"
    )


def week_summary(sales: list[Sale], days: list[date]) -> str:
    by_day: dict[date, float] = defaultdict(float)
    for s in sales:
        by_day[s.created_at.date()] += s.amount
    lines = [
        f"{WEEKDAYS[d.weekday()]} {d:%d.%m}: {money(by_day[d])}" for d in days
    ]
    total = sum(by_day.values())
    best = max(days, key=lambda d: by_day[d])
    lines.append("")
    lines.append(f"За 7 днів: <b>{money(total)}</b>")
    if total:
        lines.append(f"Найкращий день: {WEEKDAYS[best.weekday()]} {best:%d.%m}")
    return "\n".join(lines)


def build_excel(sales: list[Sale]) -> bytes:
    wb = Workbook()
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="2F5496")

    ws = wb.active
    ws.title = "Продажі"
    ws.append(["Дата", "Час", "Сума, ₴", "Коментар"])
    for s in sales:
        ws.append([s.created_at.date(), s.created_at.strftime("%H:%M"), s.amount, s.note])

    by_day: dict[date, list[float]] = defaultdict(list)
    for s in sales:
        by_day[s.created_at.date()].append(s.amount)
    days = wb.create_sheet("По днях")
    days.append(["Дата", "Продажів", "Виручка, ₴", "Середній чек, ₴"])
    for d in sorted(by_day):
        amounts = by_day[d]
        days.append([d, len(amounts), sum(amounts), round(sum(amounts) / len(amounts), 2)])

    for sheet in (ws, days):
        for cell in sheet[1]:
            cell.font, cell.fill = header_font, header_fill
        for row in sheet.iter_rows(min_row=2):
            row[0].number_format = "DD.MM.YYYY"
            for cell in row[2:]:
                if isinstance(cell.value, (int, float)):
                    cell.number_format = "#,##0.00"
        for i, width in enumerate((12, 10, 14, 30), start=1):
            sheet.column_dimensions[get_column_letter(i)].width = width
        sheet.freeze_panes = "A2"

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
